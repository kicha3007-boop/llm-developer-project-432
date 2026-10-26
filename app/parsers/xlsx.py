"""Разбор xlsx: каждая таблица листа со строкой заголовков колонок.

Заголовки лежат первой непустой строкой — это обязательная часть: «800» без «мощность, Вт»
ничего не значит. Резка превращает каждую строку в текст «артикул: параметр значение».
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from app.parsers import Block, ParsedDocument, ParseError


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def parse_xlsx(path: str | Path) -> ParsedDocument:
    try:
        workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    except Exception as exc:
        raise ParseError(f"xlsx не открывается: {exc}") from exc

    blocks: list[Block] = []
    title: str | None = None
    try:
        for sheet_no, sheet in enumerate(workbook.worksheets, start=1):
            rows = [[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
            rows = [r for r in rows if any(r)]
            if not rows:
                continue
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            # пустые колонки справа/между убираем по всей таблице
            keep = [i for i in range(width) if any(r[i] for r in rows)]
            rows = [[r[i] for i in keep] for r in rows]
            blocks.append(Block(kind="heading", text=sheet.title, page=sheet_no, level=1))
            blocks.append(
                Block(kind="table", text="", page=sheet_no, rows=rows, meta={"sheet": sheet.title})
            )
            if title is None:
                title = sheet.title
    finally:
        workbook.close()
    if not blocks:
        raise ParseError("в xlsx нет непустых листов")
    return ParsedDocument(blocks=blocks, pages=None, title=title, fmt="xlsx")
